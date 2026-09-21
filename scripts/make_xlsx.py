# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
import datetime as dt

import sys
OUT = sys.argv[1] if len(sys.argv) > 1 else 'career-tracker.xlsx'
F = 'Arial'
HDR = PatternFill('solid', fgColor='1F3864'); HF = Font(name=F, bold=True, color='FFFFFF', size=11)
thin = Side(style='thin', color='BFBFBF'); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
STATUS = ['待投递','待提交','已投递','简历筛选中','笔试/测评','一面','二面','三面/终面','HR面','Offer','未通过','已放弃']
PRIO = ['高','中','低']
ROWS = 60

def build(ws, headers, widths, rows, status_col, prio_col, date_cols, type_col=None, type_list=None, legend=''):
    ws['A1'] = legend; ws['A1'].font = Font(name=F, italic=True, color='595959', size=10)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers)); ws.row_dimensions[1].height = 32
    ws['A1'].alignment = Alignment(wrap_text=True, vertical='center')
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=c, value=h); cell.fill = HDR; cell.font = HF; cell.border = BORDER
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = widths[c-1]
    ws.row_dimensions[2].height = 30
    for r, row in enumerate(rows, 3):
        for c, v in enumerate(row, 1): ws.cell(row=r, column=c, value=v)
    for r in range(3, 3 + ROWS):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c); cell.font = Font(name=F, size=10); cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if c in date_cols: cell.number_format = 'yyyy-mm-dd'; cell.alignment = Alignment(horizontal='center', vertical='top')
    rng = lambda col: f'{get_column_letter(col)}3:{get_column_letter(col)}{2+ROWS}'
    dv = DataValidation(type='list', formula1='"' + ','.join(STATUS) + '"', allow_blank=True); ws.add_data_validation(dv); dv.add(rng(status_col))
    dp = DataValidation(type='list', formula1='"' + ','.join(PRIO) + '"', allow_blank=True); ws.add_data_validation(dp); dp.add(rng(prio_col))
    if type_col:
        dt_ = DataValidation(type='list', formula1='"' + ','.join(type_list) + '"', allow_blank=True); ws.add_data_validation(dt_); dt_.add(rng(type_col))
    sc = get_column_letter(status_col); full = f'A3:{get_column_letter(len(headers))}{2+ROWS}'
    ws.conditional_formatting.add(full, FormulaRule(formula=[f'${sc}3="Offer"'], fill=PatternFill('solid', fgColor='C6EFCE')))
    ws.conditional_formatting.add(full, FormulaRule(formula=[f'OR(${sc}3="未通过",${sc}3="已放弃")'], fill=PatternFill('solid', fgColor='E7E6E6')))
    ws.conditional_formatting.add(full, FormulaRule(formula=[f'OR(${sc}3="一面",${sc}3="二面",${sc}3="三面/终面",${sc}3="HR面",${sc}3="笔试/测评")'], fill=PatternFill('solid', fgColor='FFF2CC')))
    ws.conditional_formatting.add(full, FormulaRule(formula=[f'OR(${sc}3="待投递",${sc}3="待提交")'], fill=PatternFill('solid', fgColor='DDEBF7')))
    ws.freeze_panes = 'C3'; ws.auto_filter.ref = f'A2:{get_column_letter(len(headers))}{2+ROWS}'

wb = Workbook()
ov = wb.active; ov.title = '总览'

# ---------- 企业秋招 ----------
ws = wb.create_sheet('企业秋招')
H = ['公司','岗位名称','岗位类型','方向/意向','工作地点','优先级','投递渠道/链接','投递日期','当前状态','推进记录（时间+事件）','下一步 / 截止日期','联系人/内推','备注']
W = [16,30,10,22,12,8,38,12,12,40,22,14,40]
d = dt.date(2026, 9, 1)
R = [
 ['示例科技','【校招】规划算法工程师','校招','自动驾驶规划','深圳','高','https://example.com/jobs/1',d,'一面','2026-09-01 官网投递；2026-09-10 一面，45 分钟','等二面通知','','示例行，可删除'],
]
build(ws, H, W, R, status_col=9, prio_col=6, date_cols=[8], type_col=3, type_list=['校招','实习','博士专项','社招'],
      legend='填写说明：每个岗位一行；"岗位类型 / 优先级 / 当前状态"是下拉菜单；每次有进展，在"推进记录"里追加一行"日期 + 事件"，并更新"当前状态"。整行颜色：蓝=待投，黄=面试中，绿=Offer，灰=结束。第一行是示例，可删除。')

# ---------- 博后 ----------
ws2 = wb.create_sheet('博后申请')
H2 = ['单位 / 学校','院系 / 实验室','合作导师','研究方向','地点','优先级','申请渠道 / 链接','所需材料','申请日期','当前状态','推进记录（时间+事件）','下一步 / 截止日期','备注']
W2 = [22,22,14,24,10,8,32,30,12,12,40,22,36]
R2 = [
 ['示例大学','智能系统实验室','张教授','世界模型','杭州','中','','研究计划 + CV + 推荐信',None,'待投递','','准备研究计划','示例行，可删除'],
]
build(ws2, H2, W2, R2, status_col=10, prio_col=6, date_cols=[9],
      legend='填写说明：每个博后机会一行。"当前状态"下拉可选；面试/答辩、导师沟通、材料补交等都记在"推进记录"。')

# ---------- 高校 ----------
ws3 = wb.create_sheet('高校教职')
H3 = ['学校','学院 / 部门','岗位类型','学科方向','地点','优先级','招聘公告 / 链接','所需材料','申请日期','当前状态','推进记录（时间+事件）','下一步 / 截止日期','备注']
W3 = [20,20,14,20,10,8,32,30,12,12,40,22,36]
R3 = [
 ['示例学院','交通学院','教研岗','智能交通','上海','中','','CV + 教学陈述 + 研究陈述',None,'待投递','','','示例行，可删除'],
]
build(ws3, H3, W3, R3, status_col=10, prio_col=6, date_cols=[9], type_col=3, type_list=['教研岗','高层次人才（教师）','科研岗','科研行政岗','其他'],
      legend='填写说明：每个高校岗位一行。"岗位类型"可下拉选择，也可直接改成其他文字。试讲、学院面试、校级审批等节点记在"推进记录"。')

# ---------- 总览 ----------
ov['A1'] = '秋招进度总览'; ov['A1'].font = Font(name=F, bold=True, size=16, color='1F3864')
ov['A2'] = '本页全部是公式，自动统计另外三个表；不需要手动填写。'; ov['A2'].font = Font(name=F, italic=True, color='595959', size=10)
heads = ['状态','企业秋招','博后申请','高校教职','合计']
for c, h in enumerate(heads, 1):
    cell = ov.cell(row=4, column=c, value=h); cell.fill = HDR; cell.font = HF; cell.border = BORDER; cell.alignment = Alignment(horizontal='center')
for i, s in enumerate(STATUS):
    r = 5 + i
    ov.cell(row=r, column=1, value=s)
    ov.cell(row=r, column=2, value=f"=COUNTIF('企业秋招'!$I$3:$I${2+ROWS},A{r})")
    ov.cell(row=r, column=3, value=f"=COUNTIF('博后申请'!$J$3:$J${2+ROWS},A{r})")
    ov.cell(row=r, column=4, value=f"=COUNTIF('高校教职'!$J$3:$J${2+ROWS},A{r})")
    ov.cell(row=r, column=5, value=f'=SUM(B{r}:D{r})')
tr = 5 + len(STATUS)
ov.cell(row=tr, column=1, value='已记录条目数')
ov.cell(row=tr, column=2, value=f"=COUNTA('企业秋招'!$A$3:$A${2+ROWS})")
ov.cell(row=tr, column=3, value=f"=COUNTA('博后申请'!$A$3:$A${2+ROWS})")
ov.cell(row=tr, column=4, value=f"=COUNTA('高校教职'!$A$3:$A${2+ROWS})")
ov.cell(row=tr, column=5, value=f'=SUM(B{tr}:D{tr})')
ov.cell(row=tr+1, column=1, value='其中：未填状态')
for c in (2,3,4):
    L = get_column_letter(c); ov.cell(row=tr+1, column=c, value=f'={L}{tr}-SUM({L}5:{L}{tr-1})')
ov.cell(row=tr+1, column=5, value=f'=SUM(B{tr+1}:D{tr+1})')
for r in range(5, tr + 2):
    for c in range(1, 6):
        cell = ov.cell(row=r, column=c); cell.font = Font(name=F, size=11, bold=(r >= tr)); cell.border = BORDER
        if c > 1: cell.alignment = Alignment(horizontal='center')
for c, w in enumerate([20,12,12,12,10], 1): ov.column_dimensions[get_column_letter(c)].width = w
ov.cell(row=tr+3, column=1, value='状态口径').font = Font(name=F, bold=True, size=11)
notes = ['待投递：已选定但还没开始填；待提交：表单已填好但未点提交','已投递 → 简历筛选中 → 笔试/测评 → 一面 → 二面 → 三面/终面 → HR面 → Offer','未通过：任一环节被拒；已放弃：自己主动不再推进','需要新增状态时，修改三个表里"当前状态"列的数据验证列表，并在本页 A 列加一行']
for i, n in enumerate(notes):
    ov.cell(row=tr+4+i, column=1, value=n).font = Font(name=F, size=10, color='595959')
wb.save(OUT); print('saved', OUT)
